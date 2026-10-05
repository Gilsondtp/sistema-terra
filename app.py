import io
import os
import sys
import re
import json
import shutil
import zipfile
from datetime import date, datetime, timedelta

# Quando executado com pythonw.exe no Windows (sem janela de terminal),
# redireciona stdout/stderr para um arquivo de log para evitar erros de console
if sys.stdout is None or sys.stderr is None:
    _log_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sistema_terra.log")
    _log_file = open(_log_path, "a", encoding="utf-8", buffering=1)
    if sys.stdout is None:
        sys.stdout = _log_file
    if sys.stderr is None:
        sys.stderr = _log_file
from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    jsonify,
    send_file,
    Response,
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import database
import hotfolder_monitor

app = Flask(__name__)
app.secret_key = "sistema_terra_secret_key_sqlite_python"

# Inicializa banco SQLite e inicia thread de monitoramento Hot Folder
database.init_db()
watcher_thread = hotfolder_monitor.HotFolderWatcherThread()
watcher_thread.start()


# Filtros Jinja2 para formatação brasileira
@app.template_filter("moeda")
def filter_moeda(val):
    return database.format_money_br(val)


@app.template_filter("data_br")
def filter_data_br(val):
    return database.format_date_br(val)


@app.template_filter("data_curta")
def filter_data_curta(val):
    return database.format_date_short(val)


@app.template_filter("cnpj")
def filter_cnpj(val):
    return database.formatar_cnpj(val)


@app.context_processor
def inject_globals():
    cfg = database.get_config()
    # Conta quantos lançamentos não identificados existem no total
    conn = database.get_db_connection()
    nao_id_count = conn.execute(
        "SELECT COUNT(*) FROM lancamentos WHERE cliente_id IS NULL OR identificado = 0"
    ).fetchone()[0]
    conn.close()
    return {
        "site_title": "Sistema Terra",
        "cfg": cfg,
        "global_nao_identificados": nao_id_count,
        "hoje_iso": date.today().isoformat(),
        "hoje_br": date.today().strftime("%d/%m/%Y"),
    }


def _extrair_servicos_form(req):
    descricoes = req.form.getlist("servico_descricao[]")
    quantidades = req.form.getlist("servico_quantidade[]")
    produtos_ids = req.form.getlist("servico_produto[]")
    valores = req.form.getlist("servico_valor[]")
    cores_list = req.form.getlist("servico_cores[]")
    pastas_list = req.form.getlist("servico_pasta[]")

    servicos = []
    for i in range(len(descricoes)):
        desc = (descricoes[i] or "").strip()
        qtd_raw = quantidades[i] if i < len(quantidades) else "0"
        prod_raw = produtos_ids[i] if i < len(produtos_ids) else ""
        val_raw = valores[i] if i < len(valores) else "0,00"
        cor_raw = cores_list[i] if i < len(cores_list) else ""
        pasta_raw = pastas_list[i] if i < len(pastas_list) else ""

        try:
            qtd = int(qtd_raw or 0)
        except ValueError:
            qtd = 0

        # Permite criar serviço se tiver descrição ou produto ou valor > 0
        if desc or prod_raw or database.parse_money(val_raw) > 0:
            servicos.append(
                {
                    "descricao": desc or "Serviço",
                    "quantidade": qtd,
                    "produto_id": prod_raw if prod_raw else None,
                    "valor": val_raw,
                    "cores_detectadas": cor_raw,
                    "pasta_origem": pasta_raw,
                }
            )
    return servicos


def _resumir_arquivos_origem(servico):
    """Formata os nomes TIFF de origem para consulta no tooltip do serviço diário."""
    origem = servico.get("arquivos_origem") or ""
    if isinstance(origem, str):
        try:
            origem = json.loads(origem)
        except (TypeError, ValueError):
            origem = [origem] if origem else []
    if not isinstance(origem, (list, tuple)):
        return ""
    nomes = [str(nome) for nome in origem if nome]
    if len(nomes) > 4:
        return ", ".join(nomes[:4]) + f" e mais {len(nomes) - 4} arquivo(s)"
    return ", ".join(nomes)


# ==============================================================================
# 1. TELA PRINCIPAL: LANC. DIÁRIO (PRIMEIRA ALTERAÇÃO)
# ==============================================================================
@app.route("/", methods=["GET"])
@app.route("/lancamento-diario", methods=["GET"])
def lancamento_diario():
    # Executa uma leitura rápida do Hot Folder ao abrir a tela
    novos_detectados = hotfolder_monitor.executar_ciclo_monitoramento()
    if novos_detectados:
        flash(
            f"{len(novos_detectados)} novo(s) arquivo(s) processado(s) automaticamente do Hot Folder!",
            "info",
        )

    modo_filtro = request.args.get("modo", "diario")
    # Se 'data' foi passado na URL (mesmo vazio), respeita o valor; se não foi passado, assume a data de hoje no modo diário
    if "data" in request.args:
        data_sel = request.args.get("data", "").strip()
    else:
        data_sel = date.today().isoformat() if modo_filtro == "diario" else ""

    data_inicio = request.args.get("data_inicio", "").strip()
    data_fim = request.args.get("data_fim", "").strip()
    cliente_id = request.args.get("cliente_id", "").strip()
    produto_id = request.args.get("produto_id", "").strip()
    texto_servico = request.args.get("texto_servico", "").strip()

    filtros = {"order_by": "l.identificado ASC, l.data_lancamento DESC, l.id DESC"}

    if data_inicio and data_fim:
        filtros["data_inicio"] = data_inicio
        filtros["data_fim"] = data_fim
    elif data_sel:
        filtros["data_exata"] = data_sel

    if modo_filtro == "nao_identificados":
        filtros["cliente_id"] = "nao_identificado"
    elif cliente_id:
        filtros["cliente_id"] = cliente_id

    if produto_id:
        filtros["produto_id"] = produto_id

    if texto_servico:
        filtros["texto_servico"] = texto_servico

    lancamentos = database.listar_lancamentos(filtros)

    # Achata os serviços para a listagem diária detalhada (ou exibe linha de lançamento sem serviço se for só pgto/entrega)
    itens_diarios = []
    soma_quantidade_produtos = 0
    soma_valores = 0.0
    total_servicos_count = 0
    total_nao_identificados = 0

    for l in lancamentos:
        is_identificado = bool(l["cliente_id"] and l["identificado"])
        if not is_identificado:
            total_nao_identificados += max(1, len(l["servicos"]))

        if l["servicos"]:
            for idx, s in enumerate(l["servicos"]):
                # Se filtrou por produto específico, mostra apenas os serviços daquele produto
                if produto_id and str(s.get("produto_id") or "") != str(produto_id):
                    continue
                # Se filtrou por texto de serviço, filtra os serviços correspondentes
                if (
                    texto_servico
                    and texto_servico.lower() not in (s.get("descricao") or "").lower()
                    and texto_servico.lower() not in (l.get("cliente_nome_original") or "").lower()
                ):
                    continue

                qtd = int(s.get("quantidade") or 0)
                val = float(s.get("valor") or 0.0)
                soma_quantidade_produtos += qtd
                soma_valores += val
                total_servicos_count += 1

                itens_diarios.append(
                    {
                        "lancamento_id": l["id"],
                        "servico_id": s["id"],
                        "data_lancamento": l["data_lancamento"],
                        "cliente_id": l["cliente_id"],
                        "cliente_nome": l["cliente_nome"],
                        "cliente_nome_original": l.get("cliente_nome_original") or "",
                        "identificado": is_identificado,
                        "origem": l.get("origem") or "manual",
                        "descricao": s.get("descricao") or "",
                        "quantidade": qtd,
                        "produto_id": s.get("produto_id"),
                        "produto_nome": s.get("produto_nome") or "-",
                        "pasta_origem": s.get("pasta_origem") or "",
                        "cores_detectadas": s.get("cores_detectadas") or "",
                        "arquivos_origem": _resumir_arquivos_origem(s),
                        "valor": val,
                        "valor_entrega": float(l["valor_entrega"] or 0.0) if idx == 0 else 0.0,
                        "valor_pagamento": float(l["valor_pagamento"] or 0.0) if idx == 0 else 0.0,
                        "lancamento_obj": l,
                    }
                )
        else:
            total_servicos_count += 1
            itens_diarios.append(
                {
                    "lancamento_id": l["id"],
                    "servico_id": None,
                    "data_lancamento": l["data_lancamento"],
                    "cliente_id": l["cliente_id"],
                    "cliente_nome": l["cliente_nome"],
                    "cliente_nome_original": l.get("cliente_nome_original") or "",
                    "identificado": is_identificado,
                    "origem": l.get("origem") or "manual",
                    "descricao": l.get("observacao") or "(Pagamento / Entrega avulsa)",
                    "quantidade": 0,
                    "produto_id": None,
                    "produto_nome": "-",
                    "pasta_origem": "",
                    "cores_detectadas": "",
                    "arquivos_origem": "",
                    "valor": 0.0,
                    "valor_entrega": float(l["valor_entrega"] or 0.0),
                    "valor_pagamento": float(l["valor_pagamento"] or 0.0),
                    "lancamento_obj": l,
                }
            )

    clientes = database.listar_clientes()
    produtos = database.listar_produtos()

    return render_template(
        "lancamento_diario.html",
        active_page="diario",
        itens_diarios=itens_diarios,
        clientes=clientes,
        produtos=produtos,
        modo_filtro=modo_filtro,
        data_sel=data_sel,
        data_inicio=data_inicio,
        data_fim=data_fim,
        cliente_id=str(cliente_id),
        produto_id=str(produto_id),
        texto_servico=texto_servico,
        soma_quantidade_produtos=soma_quantidade_produtos,
        soma_valores=soma_valores,
        total_servicos_count=total_servicos_count,
        total_nao_identificados=total_nao_identificados,
        data_corte_default=(date.today().replace(day=1) - timedelta(days=1)).isoformat(),
    )


@app.route("/lancamento/salvar", methods=["POST"])
def salvar_lancamento():
    acao = request.form.get("acao", "criar")
    redirect_to = request.form.get("redirect_to") or url_for("lancamento_diario")
    try:
        dados = {
            "cliente_id": request.form.get("cliente_id"),
            "cliente_nome_original": request.form.get("cliente_nome_original", ""),
            "data_lancamento": request.form.get("data_lancamento"),
            "valor_entrega": request.form.get("valor_entrega", "0,00"),
            "valor_pagamento": request.form.get("valor_pagamento", "0,00"),
            "fatura_anterior": request.form.get("fatura_anterior", "0,00"),
            "pagamento_fatura_anterior": request.form.get("pagamento_fatura_anterior", "0,00"),
            "observacao": request.form.get("observacao", ""),
            "servicos": _extrair_servicos_form(request),
        }
        if acao == "atualizar" and request.form.get("id"):
            database.atualizar_lancamento(request.form.get("id"), dados)
            flash("Lançamento atualizado com sucesso!", "success")
        else:
            database.criar_lancamento(dados)
            flash("Lançamento registrado com sucesso!", "success")
    except Exception as e:
        flash(f"Erro ao salvar lançamento: {e}", "danger")
    return redirect(redirect_to)


@app.route("/lancamento/excluir", methods=["POST"])
def excluir_lancamento_route():
    lanc_id = request.form.get("id")
    servico_id = request.form.get("servico_id")
    redirect_to = request.form.get("redirect_to") or url_for("lancamento_diario")
    if servico_id and str(servico_id).isdigit() and int(servico_id) > 0:
        ok = database.excluir_servico_diario(int(servico_id), lanc_id)
    elif lanc_id:
        ok = database.excluir_lancamento(lanc_id)
    else:
        ok = False

    if ok:
        flash("Registro excluído com sucesso e estoque atualizado!", "success")
    else:
        flash("Erro ao excluir registro.", "danger")
    return redirect(redirect_to)


@app.route("/virada-mes", methods=["POST"])
def virada_mes_route():
    data_corte = request.form.get("data_corte") or (date.today().replace(day=1) - timedelta(days=1)).isoformat()
    modo_limpeza = request.form.get("modo_limpeza") or "somente_hotfolder_mes_atual"
    redirect_to = request.form.get("redirect_to") or url_for("lancamento_diario")
    ok, msg = database.realizar_virada_mes(data_corte=data_corte, modo_limpeza=modo_limpeza)
    if ok:
        msg += " Os lançamentos do Caixa foram preservados e não foram apagados."
    flash(msg, "success" if ok else "danger")
    return redirect(redirect_to)


@app.route("/lancamento/atribuir-cliente", methods=["POST"])
def atribuir_cliente_route():
    lanc_id = request.form.get("lancamento_id")
    cli_id = request.form.get("cliente_id")
    salvar_apelido = request.form.get("salvar_apelido") == "1"
    redirect_to = request.form.get("redirect_to") or url_for("lancamento_diario")

    if not lanc_id or not cli_id:
        flash("Selecione um cliente válido para atribuir ao serviço.", "warning")
        return redirect(redirect_to)

    if database.atribuir_cliente_lancamento(lanc_id, cli_id, salvar_apelido=salvar_apelido):
        flash("Cliente atribuído ao lançamento com sucesso!", "success")
    else:
        flash("Não foi possível atribuir o cliente.", "danger")
    return redirect(redirect_to)


# ==============================================================================
# 2. TELA LANÇAMENTOS (COMPLETA ORIGINAL PRESERVADA)
# ==============================================================================
@app.route("/lancamentos", methods=["GET"])
def lancamentos_page():
    filtros = {}
    data_inicio = request.args.get("data_inicio", "")
    data_fim = request.args.get("data_fim", "")
    cliente_id = request.args.get("cliente_id", "")
    produto_id = request.args.get("produto_id", "")
    texto_servico = request.args.get("texto_servico", "").strip()

    if data_inicio:
        filtros["data_inicio"] = data_inicio
    if data_fim:
        filtros["data_fim"] = data_fim
    if cliente_id:
        filtros["cliente_id"] = cliente_id
    if produto_id:
        filtros["produto_id"] = produto_id
    if texto_servico:
        filtros["texto_servico"] = texto_servico

    filtros["order_by"] = "l.data_lancamento DESC, l.id DESC"
    lancamentos = database.listar_lancamentos(filtros)
    clientes = database.listar_clientes()
    produtos = database.listar_produtos()

    return render_template(
        "lancamentos.html",
        active_page="lancamentos",
        lancamentos=lancamentos,
        clientes=clientes,
        produtos=produtos,
        data_inicio=data_inicio,
        data_fim=data_fim,
        cliente_id=str(cliente_id),
        produto_id=str(produto_id),
        texto_servico=texto_servico,
    )


# ==============================================================================
# 3. TELA CLIENTES
# ==============================================================================
@app.route("/clientes", methods=["GET", "POST"])
def clientes_page():
    if request.method == "POST":
        acao = request.form.get("acao")
        if acao in ("criar", "atualizar"):
            cid = request.form.get("id")
            nome = (request.form.get("nome") or "").strip()
            telefone = request.form.get("telefone", "")
            rua_bairro = request.form.get("rua_bairro", "")
            cidade = request.form.get("cidade", "")
            fatura_anterior = request.form.get("fatura_anterior", "0,00")
            apelidos = request.form.get("apelidos", "")
            cnpj = request.form.get("cnpj", "")

            if not nome:
                flash("O nome do cliente é obrigatório!", "danger")
            else:
                if acao == "atualizar" and cid:
                    database.atualizar_cliente(
                        cid, nome, telefone, rua_bairro, cidade, fatura_anterior, apelidos, cnpj
                    )
                    flash("Cliente atualizado com sucesso!", "success")
                else:
                    database.criar_cliente(
                        nome, telefone, rua_bairro, cidade, fatura_anterior, apelidos, cnpj
                    )
                    flash("Cliente cadastrado com sucesso!", "success")
        elif acao == "excluir":
            cid = request.form.get("id")
            ok, msg = database.excluir_cliente(cid)
            flash(msg, "success" if ok else "danger")
        return redirect(url_for("clientes_page"))

    clientes = database.listar_clientes()
    return render_template("clientes.html", active_page="clientes", clientes=clientes)


# ==============================================================================
# 4. TELA PRODUTOS & ESTOQUE INTEGRADOS (SEGUNDA ALTERAÇÃO)
# ==============================================================================
@app.route("/produtos", methods=["GET", "POST"])
@app.route("/estoque", methods=["GET", "POST"])
def produtos_page():
    if request.method == "POST":
        acao = request.form.get("acao")
        if acao in ("criar", "atualizar"):
            pid = request.form.get("id")
            nome = (request.form.get("nome") or "").strip()
            pasta_hotfolder = (request.form.get("pasta_hotfolder") or "").strip()
            estoque_inicial = request.form.get("estoque_inicial", 0)
            estoque_atual_override = request.form.get("estoque_atual", "")
            data_cadastro = request.form.get("data_cadastro") or date.today().isoformat()
            preco = request.form.get("preco", "0,00")
            valor_compra = request.form.get("valor_compra", "0,00")
            anotacoes = request.form.get("anotacoes", "")

            if not nome:
                flash("O nome do produto (chapa) é obrigatório!", "danger")
            else:
                if acao == "atualizar" and pid:
                    database.atualizar_produto(
                        pid,
                        nome,
                        estoque_inicial,
                        data_cadastro,
                        preco,
                        valor_compra,
                        anotacoes,
                        pasta_hotfolder=pasta_hotfolder,
                        estoque_atual_override=estoque_atual_override,
                    )
                    flash("Produto / Estoque atualizado com sucesso!", "success")
                else:
                    database.criar_produto(
                        nome,
                        estoque_inicial,
                        data_cadastro,
                        preco,
                        valor_compra,
                        anotacoes,
                        pasta_hotfolder=pasta_hotfolder,
                    )
                    flash("Produto cadastrado com sucesso!", "success")
        elif acao == "movimentar_estoque":
            pid = request.form.get("id")
            tipo = request.form.get("tipo_movimento", "entrada")
            qtd = request.form.get("quantidade", 0)
            if database.movimentar_estoque_produto(pid, qtd, tipo):
                flash("Estoque movimentado com sucesso!", "success")
            else:
                flash("Erro ao movimentar estoque.", "danger")
        elif acao == "virar_mes_todos":
            for p in database.listar_produtos():
                database.movimentar_estoque_produto(p["id"], 0, "virar_mes")
            flash("Estoque inicial de todos os produtos igualado ao estoque atual para o novo mês!", "success")
        elif acao == "excluir":
            pid = request.form.get("id")
            ok, msg = database.excluir_produto(pid)
            flash(msg, "success" if ok else "danger")
        return redirect(url_for("produtos_page"))

    produtos = database.listar_produtos()
    totais_estoque = {
        "estoque_inicial": sum(p["estoque_inicial"] for p in produtos),
        "estoque_atual": sum(p["estoque_atual"] for p in produtos),
        "consumido": sum(p["estoque_consumido"] for p in produtos),
        "valor_estoque_inicial": sum(p["valor_estoque_inicial"] for p in produtos),
        "valor_estoque_atual": sum(p["valor_estoque_atual"] for p in produtos),
        "custo_consumido": sum(p["custo_chapas_consumidas"] for p in produtos),
        "media_compra": (sum(float(p["valor_compra"] or 0) for p in produtos) / len(produtos)) if produtos else 0.0,
        "media_venda": (sum(float(p["preco"] or 0) for p in produtos) / len(produtos)) if produtos else 0.0,
    }
    return render_template(
        "produtos.html",
        active_page="produtos",
        produtos=produtos,
        totais_estoque=totais_estoque,
    )


# ==============================================================================
# 5. TELA RELATÓRIO (FATURA INDIVIDUAL)
# ==============================================================================
@app.route("/relatorios", methods=["GET", "POST"])
def relatorios_page():
    if request.method == "POST" and request.form.get("fechar_fatura") == "1":
        cliente_id = request.form.get("cliente_id")
        data_inicio = request.form.get("data_inicio") or None
        data_fim = request.form.get("data_fim") or None
        if not cliente_id:
            flash("Selecione um cliente para fechar a fatura.", "danger")
            return redirect(url_for("relatorios_page"))
        ok, resultado = database.fechar_fatura_cliente(cliente_id, data_inicio, data_fim)
        if ok:
            flash(
                f"Fatura fechada com sucesso! Saldo devedor de R$ {database.format_money_br(resultado)} "
                f"transferido para a Fatura Anterior do cliente.",
                "success",
            )
        else:
            flash(f"Erro ao fechar fatura: {resultado}", "danger")
        return redirect(url_for("relatorios_page", cliente_id=cliente_id))

    clientes = database.listar_clientes()
    cliente_id = request.values.get("cliente_id", "")
    data_inicio = request.values.get("data_inicio", "")
    data_fim = request.values.get("data_fim", "")
    # Ao escolher um cliente (ou alterar o período), a fatura é exibida automaticamente.
    fatura_gerada = bool(cliente_id)

    lancamentos = []
    cliente_atual = None
    totais = {
        "servicos": 0.0,
        "entrega": 0.0,
        "pagamento": 0.0,
        "fatura_anterior": 0.0,
        "pagamento_fatura_anterior": 0.0,
        "saldo_fatura_anterior": 0.0,
        "total_fatura": 0.0,
        "total_pago": 0.0,
        "saldo_devedor": 0.0,
        "quantidade": 0,
    }
    produtos_utilizados = {}
    observacoes_relatorio = []
    primeira_data = None
    ultima_data = None

    if cliente_id:
        cliente_atual = database.buscar_cliente(int(cliente_id))
        filtros = {"cliente_id": int(cliente_id)}
        if data_inicio:
            filtros["data_inicio"] = data_inicio
        if data_fim:
            filtros["data_fim"] = data_fim
        lancamentos = database.listar_lancamentos(filtros)

        fatura_anterior_cli = float(cliente_atual["fatura_anterior"] or 0.0) if cliente_atual else 0.0
        pgto_fatura_ant_soma = 0.0

        for l in lancamentos:
            dt_l = l["data_lancamento"]
            if primeira_data is None or dt_l < primeira_data:
                primeira_data = dt_l
            if ultima_data is None or dt_l > ultima_data:
                ultima_data = dt_l

            if l.get("observacao") and not l["observacao"].startswith("Auto Hot Folder"):
                if l["observacao"] not in observacoes_relatorio:
                    observacoes_relatorio.append(l["observacao"])

            for s in l["servicos"]:
                totais["servicos"] += float(s["valor"] or 0.0)
                totais["quantidade"] += int(s["quantidade"] or 0)
                pnome = s.get("produto_nome")
                if pnome:
                    produtos_utilizados[pnome] = produtos_utilizados.get(pnome, 0) + int(s["quantidade"] or 0)

            totais["entrega"] += float(l["valor_entrega"] or 0.0)
            totais["pagamento"] += float(l["valor_pagamento"] or 0.0)
            pgto_fatura_ant_soma += float(l["pagamento_fatura_anterior"] or 0.0)

        totais["fatura_anterior"] = fatura_anterior_cli
        totais["pagamento_fatura_anterior"] = pgto_fatura_ant_soma
        totais["saldo_fatura_anterior"] = fatura_anterior_cli - pgto_fatura_ant_soma
        totais["total_fatura"] = totais["servicos"] + totais["entrega"]
        totais["total_pago"] = totais["pagamento"] + pgto_fatura_ant_soma
        totais["saldo_devedor"] = (
            totais["servicos"] + totais["entrega"] - totais["pagamento"]
        ) + totais["saldo_fatura_anterior"]

    produtos_utilizados_ordenados = sorted(produtos_utilizados.items(), key=lambda x: x[0].lower())

    return render_template(
        "relatorios.html",
        active_page="relatorios",
        clientes=clientes,
        cliente_id=str(cliente_id),
        cliente_atual=cliente_atual,
        data_inicio=data_inicio,
        data_fim=data_fim,
        lancamentos=lancamentos,
        totais=totais,
        produtos_utilizados=produtos_utilizados_ordenados,
        observacoes_relatorio=observacoes_relatorio,
        primeira_data=primeira_data,
        ultima_data=ultima_data,
        fatura_gerada=fatura_gerada,
    )


# ==============================================================================
# 6. TELA DEMONSTRATIVO
# ==============================================================================
def _calcular_dados_demonstrativo(data_inicio, data_fim, produto_id=None):
    filtros = {
        "data_inicio": data_inicio,
        "data_fim": data_fim,
        "apenas_identificados": True,
    }
    if produto_id:
        filtros["produto_id"] = int(produto_id)

    lancamentos = database.listar_lancamentos(filtros)
    lista_clientes = database.listar_clientes()

    dados_clientes = {}
    if not produto_id:
        for c in lista_clientes:
            dados_clientes[c["id"]] = {
                "id": c["id"],
                "nome": c["nome"],
                "fatura_anterior": float(c["fatura_anterior"] or 0.0),
                "servicos": 0.0,
                "entregas": 0.0,
                "pagamentos": 0.0,
                "saldo": 0.0,
            }

    totais_gerais = {
        "servicos": 0.0,
        "entregas": 0.0,
        "pagamentos": 0.0,
        "saldo": 0.0,
    }
    produtos_utilizados = {}

    for l in lancamentos:
        cid = l["cliente_id"]
        if not cid:
            continue

        # Se há filtro de produto, verifica se o lançamento tem o produto
        if produto_id:
            usou = any(str(s.get("produto_id") or "") == str(produto_id) for s in l["servicos"])
            if not usou:
                continue

        if cid not in dados_clientes:
            dados_clientes[cid] = {
                "id": cid,
                "nome": l["cliente_nome"],
                "fatura_anterior": float(l.get("fatura_anterior") or 0.0),
                "servicos": 0.0,
                "entregas": 0.0,
                "pagamentos": 0.0,
                "saldo": 0.0,
            }

        for s in l["servicos"]:
            if produto_id and str(s.get("produto_id") or "") != str(produto_id):
                continue
            val_s = float(s["valor"] or 0.0)
            qtd_s = int(s["quantidade"] or 0)
            dados_clientes[cid]["servicos"] += val_s
            totais_gerais["servicos"] += val_s

            pnome = s.get("produto_nome")
            if pnome:
                if pnome not in produtos_utilizados:
                    produtos_utilizados[pnome] = {"quantidade": 0, "valor_total": 0.0}
                produtos_utilizados[pnome]["quantidade"] += qtd_s
                produtos_utilizados[pnome]["valor_total"] += val_s

        ent = float(l["valor_entrega"] or 0.0)
        pgto = float(l["valor_pagamento"] or 0.0) + float(l.get("pagamento_fatura_anterior") or 0.0)
        dados_clientes[cid]["entregas"] += ent
        dados_clientes[cid]["pagamentos"] += pgto
        totais_gerais["entregas"] += ent
        totais_gerais["pagamentos"] += pgto

    lista_resumo = []
    for cid, d in dados_clientes.items():
        d["saldo"] = d["servicos"] + d["entregas"] - d["pagamentos"]
        lista_resumo.append(d)

    lista_resumo.sort(key=lambda x: x["nome"].lower())
    totais_gerais["saldo"] = (
        totais_gerais["servicos"] + totais_gerais["entregas"] - totais_gerais["pagamentos"]
    )
    produtos_ordenados = sorted(produtos_utilizados.items(), key=lambda x: x[0].lower())
    estoque_critico = database.listar_estoque_critico(20)

    return lista_resumo, totais_gerais, produtos_ordenados, estoque_critico


@app.route("/demonstrativo", methods=["GET", "POST"])
def demonstrativo_page():
    hoje = date.today()
    primeiro_dia = hoje.replace(day=1).isoformat()
    # Último dia do mês
    prox_mes = (hoje.replace(day=28) + timedelta(days=4)).replace(day=1)
    ultimo_dia = (prox_mes - timedelta(days=1)).isoformat()

    data_inicio = request.values.get("data_inicio") or primeiro_dia
    data_fim = request.values.get("data_fim") or ultimo_dia
    produto_id = request.values.get("produto_id") or ""

    dados_clientes, totais_gerais, produtos_utilizados, estoque_critico = _calcular_dados_demonstrativo(
        data_inicio, data_fim, produto_id if produto_id else None
    )
    produtos = database.listar_produtos()

    return render_template(
        "demonstrativo.html",
        active_page="demonstrativo",
        data_inicio=database.parse_date_iso(data_inicio),
        data_fim=database.parse_date_iso(data_fim),
        produto_id=str(produto_id),
        produtos=produtos,
        dados_clientes=dados_clientes,
        totais_gerais=totais_gerais,
        produtos_utilizados=produtos_utilizados,
        estoque_critico=estoque_critico,
    )


@app.route("/demonstrativo/exportar-excel", methods=["GET"])
def exportar_demonstrativo_excel():
    hoje = date.today()
    primeiro_dia = hoje.replace(day=1).isoformat()
    prox_mes = (hoje.replace(day=28) + timedelta(days=4)).replace(day=1)
    ultimo_dia = (prox_mes - timedelta(days=1)).isoformat()

    data_inicio = request.args.get("data_inicio") or primeiro_dia
    data_fim = request.args.get("data_fim") or ultimo_dia
    produto_id = request.args.get("produto_id") or None

    dados_clientes, totais_gerais, produtos_utilizados, _ = _calcular_dados_demonstrativo(
        data_inicio, data_fim, produto_id
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Demonstrativo Clientes"

    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)

    ws.append([f"Demonstrativo Financeiro - {database.format_date_br(data_inicio)} a {database.format_date_br(data_fim)}"])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([])

    headers = ["Cliente", "Serviços (R$)", "Entregas (R$)", "Pagamentos (R$)", "Saldo (R$)"]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    for c in dados_clientes:
        ws.append([c["nome"], c["servicos"], c["entregas"], c["pagamentos"], c["saldo"]])

    ws.append(
        [
            "TOTAL GERAL",
            totais_gerais["servicos"],
            totais_gerais["entregas"],
            totais_gerais["pagamentos"],
            totais_gerais["saldo"],
        ]
    )
    last_row = ws.max_row
    for col_idx in range(1, 6):
        ws.cell(row=last_row, column=col_idx).font = bold_font

    # Aba 2: Produtos Utilizados
    ws2 = wb.create_sheet(title="Produtos Utilizados")
    ws2.append(["Produto (Chapa)", "Quantidade", "Valor Total (R$)"])
    for col_idx in range(1, 4):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
    for pnome, pdados in produtos_utilizados:
        ws2.append([pnome, pdados["quantidade"], pdados["valor_total"]])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"demonstrativo_{data_inicio}_a_{data_fim}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ==============================================================================
# 7. TELA CAIXA (TERCEIRA ALTERAÇÃO - INTEGRADA AO BANCO E SALDOS DE CLIENTES)
# ==============================================================================
@app.route("/caixa", methods=["GET", "POST"])
def caixa_page():
    if request.method == "POST":
        acao = request.form.get("acao")
        if acao in ("criar", "atualizar"):
            caixa_id = request.form.get("id")
            descricao = (request.form.get("descricao") or "").strip()
            data_str = request.form.get("data") or date.today().isoformat()
            valor = request.form.get("valor", "0,00")
            operacao = request.form.get("operacao", "receita")
            cliente_id = request.form.get("cliente_id") or None

            if not descricao:
                flash("Informe a descrição do lançamento de caixa.", "danger")
            else:
                if acao == "atualizar" and caixa_id:
                    database.atualizar_lancamento_caixa(
                        caixa_id, descricao, data_str, valor, operacao, cliente_id
                    )
                    flash("Lançamento de caixa atualizado!", "success")
                else:
                    database.criar_lancamento_caixa(
                        descricao, data_str, valor, operacao, cliente_id
                    )
                    flash("Lançamento registrado no caixa!", "success")
        elif acao == "excluir":
            caixa_id = request.form.get("id")
            database.excluir_lancamento_caixa(caixa_id)
            flash("Lançamento de caixa excluído!", "success")
        return redirect(url_for("caixa_page"))

    data_inicio = request.args.get("data_inicio", "")
    data_fim = request.args.get("data_fim", "")
    operacao_filtro = request.args.get("operacao", "")

    lancamentos_caixa = database.listar_lancamentos_caixa(
        data_inicio if data_inicio else None,
        data_fim if data_fim else None,
        operacao_filtro if operacao_filtro else None,
    )

    # Saldos devedores e pagamentos de todos os clientes no banco unificado
    saldos_clientes, totais_clientes = database.obter_saldos_e_pagamentos_clientes(
        data_inicio if data_inicio else None,
        data_fim if data_fim else None,
    )

    resumo_caixa = {
        "receita": 0.0,
        "despesa": 0.0,
        "pago": 0.0,
        "recebido": 0.0,
        "a_receber": 0.0,
        "transferido": 0.0,
    }
    for lc in lancamentos_caixa:
        op = lc["operacao"]
        if op == "transferencia":
            op = "transferido"
        if op in resumo_caixa:
            resumo_caixa[op] += float(lc["valor"] or 0.0)

    # Incorpora automaticamente os valores de saldo devedor dos clientes e os pagamentos lançados nos lançamentos
    # (Nota: pagamentos de lançamentos já são sincronizados em lancamentos_caixa com origem='lancamento_cliente',
    #  mas garantimos que o total de saldos devedores dos clientes e faturamento de serviços apareçam de forma integrada)
    resumo_caixa["pagamentos_clientes"] = totais_clientes["pagamentos_total"]
    resumo_caixa["saldo_devedor_clientes"] = totais_clientes["saldo_devedor"]
    resumo_caixa["servicos_faturados_clientes"] = totais_clientes["servicos"] + totais_clientes["entregas"]

    # Saldo Geral do Caixa: (Receitas + Pagamentos Recebidos) - (Despesas + Pagos)
    resumo_caixa["saldo_liquido_caixa"] = (
        resumo_caixa["receita"] + resumo_caixa["recebido"] - resumo_caixa["despesa"] - resumo_caixa["pago"]
    )

    clientes = database.listar_clientes()

    return render_template(
        "caixa.html",
        active_page="caixa",
        lancamentos_caixa=lancamentos_caixa,
        saldos_clientes=saldos_clientes,
        totais_clientes=totais_clientes,
        resumo_caixa=resumo_caixa,
        clientes=clientes,
        data_inicio=data_inicio,
        data_fim=data_fim,
        operacao_filtro=operacao_filtro,
    )


# ==============================================================================
# 8. BACKUP (EXPORTAR / IMPORTAR) & CONFIGURAÇÃO HOT FOLDER
# ==============================================================================
@app.route("/configuracoes", methods=["GET", "POST"])
def configuracoes_page():
    if request.method == "POST":
        acao = request.form.get("acao")
        if acao == "salvar_hotfolder":
            database.set_config(
                {
                    "hotfolder_network_path": request.form.get("hotfolder_network_path", "").strip(),
                    "hotfolder_local_path": request.form.get("hotfolder_local_path", "").strip(),
                    "hotfolder_ativo": "1" if request.form.get("hotfolder_ativo") == "1" else "0",
                    "hotfolder_intervalo_seg": request.form.get("hotfolder_intervalo_seg", "5"),
                    "empresa_nome": request.form.get("empresa_nome", "Terra Fotolito").strip(),
                    "empresa_endereco": request.form.get("empresa_endereco", "").strip(),
                    "empresa_telefone": request.form.get("empresa_telefone", "").strip(),
                    "empresa_email": request.form.get("empresa_email", "").strip(),
                    "empresa_pix": request.form.get("empresa_pix", "").strip(),
                }
            )
            flash("Configurações salvas com sucesso!", "success")
        elif acao == "salvar_mapeamento_pastas":
            produtos = database.listar_produtos()
            conn = database.get_db_connection()
            with conn:
                for p in produtos:
                    nova_pasta = request.form.get(f"pasta_prod_{p['id']}", "").strip()
                    conn.execute(
                        "UPDATE produtos SET pasta_hotfolder = ? WHERE id = ?",
                        (nova_pasta, p["id"]),
                    )
            conn.close()
            flash("Mapeamento de pastas Hot Folder atualizado com sucesso!", "success")
        return redirect(url_for("configuracoes_page"))

    cfg = database.get_config()
    produtos = database.listar_produtos()
    conn = database.get_db_connection()
    ultimos_arquivos = conn.execute(
        """
        SELECT a.*, l.identificado, COALESCE(c.nome, 'Não Identificado') AS cliente_nome
        FROM arquivos_monitorados a
        LEFT JOIN lancamentos l ON a.lancamento_id = l.id
        LEFT JOIN clientes c ON l.cliente_id = c.id
        ORDER BY a.id DESC LIMIT 30
        """
    ).fetchall()
    conn.close()

    return render_template(
        "configuracoes.html",
        active_page="configuracoes",
        cfg=cfg,
        produtos=produtos,
        ultimos_arquivos=[dict(r) for r in ultimos_arquivos],
    )


@app.route("/hotfolder/escanear", methods=["POST"])
def escanear_hotfolder_manual():
    resetar_ignorados = request.form.get("resetar_ignorados") == "1"
    if resetar_ignorados:
        conn = database.get_db_connection()
        with conn:
            conn.execute(
                """
                DELETE FROM arquivos_monitorados
                WHERE lancamento_id IS NULL
                   OR lancamento_id NOT IN (SELECT id FROM lancamentos)
                """
            )
        conn.close()
    novos = hotfolder_monitor.executar_ciclo_monitoramento()
    if novos:
        flash(f"{len(novos)} arquivo(s) .tif/.tff detectado(s) e lançado(s) com sucesso!", "success")
    else:
        flash("Varredura concluída: nenhum arquivo .tif/.tff novo pendente nas pastas monitoradas.", "info")
    return redirect(request.form.get("redirect_to") or url_for("lancamento_diario"))


@app.route("/hotfolder/simular", methods=["POST"])
def simular_arquivos_hotfolder():
    """
    Permite simular a entrada de arquivos gerados pelo RIP (Raster Precision Screen)
    em uma das subpastas de chapas (ex: 510x400, 660x530, 650x550, 521, 745x605),
    para testar ou demonstrar o reconhecimento automático.
    """
    pasta = (request.form.get("pasta_chapa") or "510x400").strip()
    nomes_raw = request.form.get("nomes_arquivos") or ""
    local_root = os.path.join(database.BASE_DIR, "hotfolder_output")
    pasta_destino = os.path.join(local_root, pasta)
    os.makedirs(pasta_destino, exist_ok=True)

    # Aceita separação por quebra de linha ou vírgula
    linhas = []
    for l in nomes_raw.splitlines():
        for parte in l.split(","):
            if parte.strip():
                linhas.append(parte.strip())

    criados = 0
    for nome in linhas:
        base_n, ext_n = os.path.splitext(os.path.basename(nome))
        if ext_n.lower() not in (".tif", ".tiff", ".tff"):
            nome_arq = f"{os.path.basename(nome)}.tif"
        else:
            nome_arq = os.path.basename(nome)
        caminho_arq = os.path.join(pasta_destino, nome_arq)

        # Se o arquivo já foi monitorado antes e o usuário quer testar de novo, remove do registro de monitorados
        conn = database.get_db_connection()
        with conn:
            conn.execute(
                "DELETE FROM arquivos_monitorados WHERE caminho_completo = ? OR nome_arquivo = ?",
                (os.path.abspath(caminho_arq), nome_arq),
            )
        conn.close()
        with open(caminho_arq, "w", encoding="utf-8") as f:
            f.write(f"SIMULACAO RIP CTP - {nome_arq} - {datetime.now().isoformat()}\n")
        criados += 1

    novos = hotfolder_monitor.executar_ciclo_monitoramento()
    flash(
        f"{criados} arquivo(s) .tif enviado(s) para a pasta '{pasta}' e {len(novos)} processado(s) no Lanç. Diário!",
        "success",
    )
    return redirect(request.form.get("redirect_to") or url_for("lancamento_diario"))


@app.route("/backup/exportar/<formato>", methods=["GET"])
def exportar_backup(formato):
    hoje_str = date.today().isoformat()
    if formato == "sqlite":
        return send_file(
            database.DB_PATH,
            as_attachment=True,
            download_name=f"backup_sistema_terra_{hoje_str}.db",
            mimetype="application/x-sqlite3",
        )
    elif formato == "json":
        payload = database.exportar_backup_json()
        content = json.dumps(payload, ensure_ascii=False, indent=2)
        return Response(
            content,
            mimetype="application/json",
            headers={
                "Content-Disposition": f"attachment; filename=backup_sistema_terra_{hoje_str}.json"
            },
        )
    elif formato == "sql":
        sql_text = database.exportar_backup_sql()
        return Response(
            sql_text,
            mimetype="application/sql",
            headers={
                "Content-Disposition": f"attachment; filename=backup_sistema_terra_{hoje_str}.sql"
            },
        )
    flash("Formato de exportação inválido.", "danger")
    return redirect(url_for("configuracoes_page"))


@app.route("/backup/importar", methods=["POST"])
def importar_backup():
    arq = request.files.get("arquivo_backup")
    if not arq or not arq.filename:
        flash("Selecione um arquivo de backup (.db, .json ou .sql).", "danger")
        return redirect(url_for("configuracoes_page"))

    fname = arq.filename.lower()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    try:
        if fname.endswith(".db") or fname.endswith(".sqlite") or fname.endswith(".sqlite3"):
            if os.path.exists(database.DB_PATH):
                shutil.copy2(
                    database.DB_PATH,
                    os.path.join(database.BACKUP_DIR, f"auto_pre_restore_{timestamp}.db"),
                )
            arq.save(database.DB_PATH)
            database.init_db()
            flash("Banco de dados SQLite (.db) restaurado com sucesso!", "success")
        elif fname.endswith(".json"):
            payload = json.loads(arq.read().decode("utf-8", errors="replace"))
            ok, msg = database.importar_backup_json(payload)
            flash(msg, "success" if ok else "danger")
        elif fname.endswith(".sql"):
            conteudo = arq.read().decode("utf-8", errors="replace")
            if os.path.exists(database.DB_PATH):
                shutil.copy2(
                    database.DB_PATH,
                    os.path.join(database.BACKUP_DIR, f"auto_pre_sql_{timestamp}.db"),
                )
            ok, msg = database.importar_dump_sql(conteudo)
            flash(msg, "success" if ok else "danger")
        else:
            flash("Formato não suportado. Use .db, .json ou .sql.", "warning")
    except Exception as e:
        flash(f"Erro na importação do backup: {e}", "danger")

    return redirect(url_for("configuracoes_page"))


@app.route("/download-projeto-zip", methods=["GET"])
def download_projeto_zip():
    """
    Gera e envia um arquivo .zip com todos os arquivos da versão em Python + SQLite
    prontos para rodar em qualquer computador.
    """
    mem_zip = io.BytesIO()
    arquivos_incluir = [
        ".gitignore",
        "Lancamentos 30_09_final.sql",
        "README_PYTHON.md",
        "app.py",
        "database.py",
        "hotfolder_monitor.py",
        "requirements.txt",
        "iniciar_sistema.bat",
        "iniciar_sem_janela.vbs",
        "parar_sistema.bat",
        "sistema_terra.db",
        os.path.join("backups", "setembro_completo_30_09_2026.db"),
    ]
    pastas_incluir = ["templates", "static"]

    with zipfile.ZipFile(mem_zip, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for fname in arquivos_incluir:
            fpath = os.path.join(database.BASE_DIR, fname)
            if os.path.exists(fpath):
                zf.write(fpath, arcname=os.path.join("sistema-terra", fname))

        for folder in pastas_incluir:
            folder_path = os.path.join(database.BASE_DIR, folder)
            if os.path.exists(folder_path):
                for root, _, files in os.walk(folder_path):
                    for f in files:
                        full_p = os.path.join(root, f)
                        rel_p = os.path.relpath(full_p, database.BASE_DIR)
                        zf.write(full_p, arcname=os.path.join("sistema-terra", rel_p))

    mem_zip.seek(0)
    return send_file(
        mem_zip,
        as_attachment=True,
        download_name="sistema_terra_python.zip",
        mimetype="application/zip",
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
